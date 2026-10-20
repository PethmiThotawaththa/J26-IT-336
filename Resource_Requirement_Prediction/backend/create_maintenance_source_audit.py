from pathlib import Path
from datetime import datetime, date

from openpyxl import load_workbook
import pandas as pd


ROOT = Path(r"D:\Y4S1\RP\PP1\RESOURCES")
PROJECT = Path.cwd()

SOURCE_FILE = ROOT / "Forecasted routine maintenance budget 2026.xlsx"
OUTPUT_FILE = PROJECT / "data" / "processed" / "maintenance_source_audit.csv"

DATE_FORMAT = "%d/%m/%Y %I:%M:%S %p"
DATE_FORMAT_STATUS = "CONFIRMED_DD/MM/YYYY"


def parse_inspection_date(value):
    if pd.isna(value):
        return pd.NaT

    if isinstance(value, (datetime, date)):
        return pd.to_datetime(value, errors="coerce")

    try:
        return pd.Timestamp(
            datetime.strptime(str(value).strip(), DATE_FORMAT)
        )
    except (ValueError, TypeError):
        return pd.NaT


def main():
    if not SOURCE_FILE.is_file():
        raise FileNotFoundError(
            f"Source workbook not found: {SOURCE_FILE}"
        )

    records = []

    workbook = load_workbook(
        SOURCE_FILE,
        read_only=True,
        data_only=True
    )

    try:
        if "Budget" not in workbook.sheetnames:
            raise ValueError("The workbook has no 'Budget' worksheet.")

        worksheet = workbook["Budget"]

        headers = [
            cell.value
            for cell in next(
                worksheet.iter_rows(min_row=1, max_row=1)
            )
        ]

        header_positions = {
            str(header).strip(): index
            for index, header in enumerate(headers)
            if header is not None
        }

        required_headers = [
            "RoadID",
            "Inspection date",
            "Item Code",
            "Units",
            "Quantity",
            "Total Amount (Rs.)"
        ]

        missing_headers = [
            name for name in required_headers
            if name not in header_positions
        ]

        if missing_headers:
            raise ValueError(
                f"Missing workbook columns: {missing_headers}"
            )

        for excel_row, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2
        ):
            if all(value is None for value in row):
                continue

            raw_date = row[header_positions["Inspection date"]]
            parsed_date = parse_inspection_date(raw_date)

            if isinstance(raw_date, (datetime, date)):
                raw_date_output = raw_date.isoformat()
            else:
                raw_date_output = raw_date

            records.append({
                "Excel_Row": excel_row,
                "RoadID_Raw": row[header_positions["RoadID"]],
                "Inspection_Date_Raw": raw_date_output,
                "Inspection_Date_Parsed": (
                    parsed_date.strftime("%Y-%m-%d")
                    if not pd.isna(parsed_date)
                    else None
                ),
                "Item_Code": row[header_positions["Item Code"]],
                "Units": row[header_positions["Units"]],
                "Quantity": row[header_positions["Quantity"]],
                "Total_Amount_Rs": row[
                    header_positions["Total Amount (Rs.)"]
                ],
                "Data_Status": "ESTIMATE",
                "Date_Format_Status": DATE_FORMAT_STATUS,
                "Date_Parse_Status": (
                    "PARSED"
                    if not pd.isna(parsed_date)
                    else "INVALID_OR_MISSING"
                )
            })

    finally:
        workbook.close()

    df = pd.DataFrame(records)

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUTPUT_FILE, index=False, encoding="utf-8-sig")

    parsed_dates = pd.to_datetime(
        df["Inspection_Date_Parsed"],
        errors="coerce"
    )

    print("MAINTENANCE SOURCE AUDIT COMPLETED")
    print("-" * 40)
    print(f"Output: {OUTPUT_FILE}")
    print(f"Records audited: {len(df)}")
    print(f"Date convention: {DATE_FORMAT_STATUS}")
    print(f"Parsed dates: {parsed_dates.notna().sum()}")
    print(f"Invalid or missing dates: {parsed_dates.isna().sum()}")

    if parsed_dates.notna().any():
        print(
            "Date range:",
            parsed_dates.min().date(),
            "to",
            parsed_dates.max().date()
        )

    print("Source workbook was not modified.")
    print("No resource quantities were calculated.")


if __name__ == "__main__":
    main()
from pathlib import Path
from openpyxl import load_workbook
from collections import Counter

path = Path(
    r"D:\Y4S1\RP\PP1\RESOURCES"
    r"\Forecasted routine maintenance budget 2026.xlsx"
)

wb = load_workbook(path, data_only=True)
ws = wb["Budget"]

dates = [
    ws.cell(row=row, column=4).value
    for row in range(2, ws.max_row + 1)
]

print("=== DISTINCT RAW INSPECTION DATES ===")
for value, count in Counter(dates).most_common():
    print(f"{value!r} | Records: {count}")

print("\nTotal records:", len(dates))

from pathlib import Path
from openpyxl import load_workbook

path = Path(
    r"D:\Y4S1\RP\PP1\RESOURCES"
    r"\Forecasted routine maintenance budget 2026.xlsx"
)

wb = load_workbook(path, data_only=True)

print("Sheets:", wb.sheetnames)

for ws in wb.worksheets:
    print(f"\n=== SHEET: {ws.title} ===")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 8),
                            values_only=True):
        values = [v for v in row if v is not None]
        if values:
            print(values[:10])
